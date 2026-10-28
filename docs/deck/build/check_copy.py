import re, sys
from pptx import Presentation
from pptx.util import Pt

path = sys.argv[1]
prs = Presentation(path)
dash = {chr(0x2014): "em", chr(0x2013): "en"}
rows = []
minrun_overall = 99
for i, s in enumerate(prs.slides, 1):
    words = 0
    minsz = 99
    bad = []
    for sh in s.shapes:
        if not sh.has_text_frame:
            continue
        for p in sh.text_frame.paragraphs:
            t = "".join(r.text for r in p.runs)
            words += len(t.split())
            for r in p.runs:
                if r.font.size:
                    sz = r.font.size.pt
                    minsz = min(minsz, sz)
                    if sz < 10:
                        bad.append((sz, r.text[:30]))
                for ch in r.text:
                    if ch in dash:
                        bad.append((dash[ch], r.text[:30]))
    nt = s.notes_slide.notes_text_frame.text if s.has_notes_slide else ""
    nwords = len(nt.split())
    for ch in nt:
        if ch in dash:
            bad.append((dash[ch] + "-note", nt[:30]))
    minrun_overall = min(minrun_overall, minsz) if minsz < 99 else minrun_overall
    rows.append((i, minsz, words, nwords, bad))

print("slide  minpt  words  notes  flags")
for r in rows:
    print("%5d  %5.1f  %5d  %5d  %s" % (r[0], r[1], r[2], r[3], "" if not r[4] else r[4]))
print("min run overall", minrun_overall)
print("notes range", min(r[3] for r in rows), max(r[3] for r in rows))
print("dash count", sum(len(r[4]) for r in rows))
