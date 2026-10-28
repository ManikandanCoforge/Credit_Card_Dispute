
import fitz, re
from pptx import Presentation
from pptx.util import Pt

pdf = "docs/deck/out/lumen-dispute-copilot-2026-10-29.pdf"
d = fitz.open(pdf)
print("pdf pages", d.page_count, "size", d[0].rect)
text = "".join(p.get_text() for p in d)
print("em dashes", text.count("\u2014"), "en dashes", text.count("\u2013"), "hyphen-minus count", text.count("-"))
print("has 18.2", "18.2" in text, "| has 17.58", "17.58" in text)
print("has DRAFT 66", "DRAFT 66" in text, "| has 'tool calls, cost'", "tool calls, cost" in text)

p = "docs/deck/out/lumen-dispute-copilot-2026-10-29.pptx"
pr = Presentation(p)
print("slides", len(pr.slides), "size", pr.slide_width, pr.slide_height, pr.slide_width/914400, pr.slide_height/914400)
sizes = set()
missing_notes = []
for i, s in enumerate(pr.slides, 1):
    for sh in s.shapes:
        if sh.has_text_frame:
            for para in sh.text_frame.paragraphs:
                for r in para.runs:
                    if r.text.strip():
                        sizes.add(round(r.font.size.pt,1) if r.font.size else None)
    t = s.notes_slide.notes_text_frame.text.strip() if s.has_notes_slide else ""
    if len(t) < 20: missing_notes.append(i)
print("min run pt", min(x for x in sizes if x), "distinct", sorted(x for x in sizes if x))
print("slides with thin notes", missing_notes)
