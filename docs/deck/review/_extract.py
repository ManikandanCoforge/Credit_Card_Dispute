
import io, sys
from pptx import Presentation
out = io.open("docs/deck/review/_slides.txt","w",encoding="utf-8")
prs = Presentation("docs/deck/out/lumen-dispute-copilot-2026-10-29.pptx")
out.write("slides %d\n" % len(prs.slides))
for i, s in enumerate(prs.slides, 1):
    out.write("="*70+"\n")
    out.write("SLIDE %d\n" % i)
    def walk(shapes, indent=""):
        for sh in shapes:
            try:
                if sh.shape_type == 6:
                    walk(sh.shapes, indent+"  ")
            except Exception:
                pass
            if sh.has_text_frame:
                t = sh.text_frame.text.strip()
                if t:
                    out.write(indent + "[TXT] " + t.replace("\n"," | ") + "\n")
            if getattr(sh, "has_table", False):
                for r in sh.table.rows:
                    out.write(indent+"[TBL] " + " || ".join(c.text for c in r.cells) + "\n")
    walk(s.shapes)
    if s.has_notes_slide:
        n = s.notes_slide.notes_text_frame.text.strip()
        out.write("[NOTES] " + (n if n else "(empty)") + "\n")
    else:
        out.write("[NOTES] (none)\n")
out.close()
print("ok")
